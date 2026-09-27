import json
from types import SimpleNamespace as NS

import pytest
from google.genai import types
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

import zmo_summary as summary
from zmo_state import RunStore, UnitResult

VALID = json.dumps({"summary": "Summary", "keywords": ["a", "b", "c", "d", "e"]})


def workbook_run(tmp_path, *, rows=2, **options):
    source = tmp_path / "source.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Text", "Formula", "AI Summary", "AI Keywords", "AI Status"])
    for i in range(rows):
        sheet.append([f"text {i}", "=1+1", "old", "old", "complete"])
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="00FF00")
    book.create_sheet("Other")["A1"] = "preserve"
    book.save(source)
    book.close()
    config = summary.configuration("model", worksheet="Data", column="Text", **options)
    return source, RunStore.open(tmp_path / "runs", source, config)


def test_fresh_source_status_never_skips_and_exports_literal_cells(tmp_path, monkeypatch):
    source, store = workbook_run(tmp_path)
    raw = json.dumps(
        {
            "summary": '=WEBSERVICE("https://example.invalid")',
            "keywords": ["=A1", "+a", "-b", "@c", "d"],
        }
    )
    calls = []
    monkeypatch.setattr(summary.zc, "send_text", lambda *a, **k: (calls.append(1) or raw, "ok"))
    summary.run(None, source, store)
    assert len(calls) == 2
    checked = load_workbook(store.directory / "summaries.xlsx", data_only=False)
    assert checked.sheetnames == ["Data", "Other"]
    assert checked["Other"]["A1"].value == "preserve"
    sheet = checked["Data"]
    assert sheet["B2"].value == "=1+1"
    assert sheet["A1"].fill.fgColor.rgb.endswith("00FF00")
    assert sheet["C2"].value == "old"
    assert sheet["F2"].data_type == "s" and sheet["F2"].value.startswith("=WEBSERVICE")
    assert sheet["G2"].data_type == "s"
    checked.close()
    summary.run(None, source, RunStore.load(store.directory))
    assert len(calls) == 2


@pytest.mark.parametrize("value", [None, 2, {"x": 1}, ["text"]])
def test_strict_summary_type(value):
    with pytest.raises(ValueError, match="string"):
        summary.parse_response(json.dumps({"summary": value, "keywords": list("abcde")}))


@pytest.mark.parametrize("keywords", [list("ab"), list("aaaaa"), [1, 2, 3, 4, 5], "abcde"])
def test_strict_keyword_type_and_count(keywords):
    with pytest.raises(ValueError):
        summary.parse_response(json.dumps({"summary": "s", "keywords": keywords}))


@pytest.mark.parametrize("text", ["x" * 32768, "bad\x01control"])
def test_excel_limits_rejected_without_silent_truncation(text):
    with pytest.raises(ValueError, match="Excel"):
        summary.parse_response(json.dumps({"summary": text, "keywords": list("abcde")}))


def test_missing_formula_cache_is_visible_and_does_not_call_api(tmp_path, monkeypatch):
    source, store = workbook_run(tmp_path)
    book = load_workbook(source)
    book["Data"]["A2"] = "=B2"
    book.save(source)
    book.close()
    store = RunStore.open(
        tmp_path / "runs", source, summary.configuration("model", worksheet="Data", column="Text")
    )
    assert summary.preflight(source, store.manifest["config"]["options"])[
        "missing_formula_caches"
    ] == [2]
    monkeypatch.setattr(summary.zc, "send_text", lambda *a, **k: (VALID, "ok"))
    summary.run(None, source, store)
    assert store.unit("row-2")["status"] == "needs-recalculation"
    assert store.manifest["state"] == "partial"


def test_interrupted_summary_resumes_only_unfinished_rows(tmp_path, monkeypatch):
    source, store = workbook_run(tmp_path)
    calls = []

    def send(*args, **kwargs):
        calls.append(1)
        if len(calls) == 2:
            raise KeyboardInterrupt()
        return VALID, "ok"

    monkeypatch.setattr(summary.zc, "send_text", send)
    with pytest.raises(KeyboardInterrupt):
        summary.run(None, source, store)
    assert store.complete("row-2") and store.manifest["state"] == "interrupted"
    summary.run(None, source, RunStore.load(store.directory))
    assert len(calls) == 3


def test_chunked_summary_saves_and_reuses_maps(tmp_path, monkeypatch):
    source = tmp_path / "source.txt"
    source.write_text("x" * 250)
    config = summary.configuration("model", max_chunk_chars=100)
    store = RunStore.open(tmp_path / "runs", source, config)
    calls = []

    def send(*args, **kwargs):
        calls.append(1)
        if len(calls) == 4:
            raise KeyboardInterrupt()
        return VALID, "ok"

    monkeypatch.setattr(summary.zc, "send_text", send)
    with pytest.raises(KeyboardInterrupt):
        summary.run(None, source, store)
    assert len(store.units(include_chunks=True)) == 3
    summary.run(None, source, store)
    assert len(calls) == 5 and store.complete("text")


def batch_store(tmp_path):
    source, store = workbook_run(tmp_path)
    for row in (2, 3):
        store.record(UnitResult(f"row-{row}", "batch-pending", locator={"row": row}))
    store.manifest["jobs"] = [
        {
            "submission_id": "submission",
            "name": "job",
            "state": "submitted",
            "rows": ["row-2", "row-3"],
            "cleanup_pending": ["input"],
        }
    ]
    store.save()
    return source, store


def batch_client(keys, deleted, *, state=types.JobState.JOB_STATE_SUCCEEDED):
    payload = "\n".join(
        json.dumps(
            {
                "key": key,
                "response": {
                    "candidates": [
                        {"finishReason": "STOP", "content": {"parts": [{"text": VALID}]}}
                    ]
                },
            }
        )
        for key in keys
    )
    return NS(
        batches=NS(get=lambda **kw: NS(state=state, dest=NS(file_name="result"))),
        files=NS(
            download=lambda **kw: payload.encode(), delete=lambda **kw: deleted.append(kw["name"])
        ),
    )


@pytest.mark.parametrize(
    "keys,raises", [(["row-2"], False), (["row-9"], True), (["row-2", "row-2"], True)]
)
def test_batch_validates_keys_before_mutation_and_retains_remote_results(tmp_path, keys, raises):
    source, store = batch_store(tmp_path)
    deleted = []
    client = batch_client(keys, deleted)
    if raises:
        with pytest.raises(ValueError, match="Unexpected or duplicate"):
            summary.collect_batch(client, source, store)
        assert store.unit("row-2")["status"] == "batch-pending"
    else:
        summary.collect_batch(client, source, store)
        assert store.unit("row-3")["status"] == "missing-batch-result"
        assert (store.directory / "provenance.json").is_file()
        assert store.manifest["jobs"][0]["state"] == "collected"
    assert deleted == []


def test_batch_cleanup_requires_verified_mirror_or_explicit_export(tmp_path, monkeypatch):
    source, store = batch_store(tmp_path)
    store.mirror_root = tmp_path / "drive"
    deleted = []
    client = batch_client(["row-2", "row-3"], deleted)
    monkeypatch.setattr(store, "sync", lambda: False)
    summary.collect_batch(client, source, store)
    assert deleted == []
    assert summary.cleanup_remote(client, store, exported=True)
    assert set(deleted) == {"input", "result"}


def test_active_batch_blocks_resubmit_and_sync_mode(tmp_path):
    source, store = batch_store(tmp_path)
    with pytest.raises(ValueError, match="unresolved Batch"):
        summary.submit_batch(None, source, store)
    with pytest.raises(ValueError, match="unresolved Batch"):
        summary.run(None, source, store)


def test_batch_splits_jobs_and_keeps_instructions_separate(tmp_path, monkeypatch):
    source, store = workbook_run(tmp_path, rows=3)
    calls, payloads = [], []

    def upload(client, path, **kwargs):
        payloads.append(path.read_text())
        return NS(name=f"files/{len(payloads)}")

    def create(**kwargs):
        calls.append(kwargs)
        return NS(name=f"jobs/{len(calls)}")

    monkeypatch.setattr(summary.zc, "upload_media", upload)
    summary.submit_batch(NS(batches=NS(create=create)), source, store, max_rows=1)
    assert len(calls) == 3
    item = json.loads(payloads[0])
    assert item["request"]["system_instruction"]["parts"][0]["text"] == summary.SYSTEM_INSTRUCTION
    assert "source_text_json" in item["request"]["contents"][0]["parts"][0]["text"]
    assert len({j["submission_id"] for j in store.active_jobs}) == 3


def test_ambiguous_submission_blocks_duplicate_and_reconciles(tmp_path, monkeypatch):
    source, store = workbook_run(tmp_path)
    monkeypatch.setattr(summary.zc, "upload_media", lambda *a, **k: NS(name="input"))
    client = NS(batches=NS(create=lambda **kw: (_ for _ in ()).throw(TimeoutError())))
    with pytest.raises(TimeoutError):
        summary.submit_batch(client, source, store)
    with pytest.raises(ValueError):
        summary.submit_batch(client, source, store)
    job = store.active_jobs[0]
    client.batches.list = lambda: [
        NS(display_name=job["submission_id"], name="remote", state=types.JobState.JOB_STATE_PENDING)
    ]
    summary.reconcile_jobs(client, store)
    assert store.active_jobs[0]["name"] == "remote"


def test_collection_export_failure_never_marks_collected_or_deletes(tmp_path, monkeypatch):
    source, store = batch_store(tmp_path)
    deleted = []
    monkeypatch.setattr(
        summary, "export_records", lambda *a: (_ for _ in ()).throw(OSError("disk full"))
    )
    with pytest.raises(OSError, match="disk full"):
        summary.collect_batch(batch_client(["row-2", "row-3"], deleted), source, store)
    assert store.active_jobs and store.active_jobs[0]["state"] == "JOB_STATE_SUCCEEDED"
    assert deleted == []


def test_terminal_failed_batch_marks_rows_and_allows_explicit_retry(tmp_path, monkeypatch):
    source, store = batch_store(tmp_path)
    client = batch_client([], [], state=types.JobState.JOB_STATE_FAILED)
    summary.collect_batch(client, source, store)
    assert not store.active_jobs
    assert store.unit("row-2")["status"] == "batch-failed"
    monkeypatch.setattr(summary.zc, "send_text", lambda *a, **kw: (VALID, "ok"))
    summary.run(None, source, store)
    assert store.manifest["state"] == "complete"


def test_ambiguous_job_attachment_must_match_submission(tmp_path):
    source, store = batch_store(tmp_path)
    job = store.manifest["jobs"][0]
    del job["name"]
    client = NS(batches=NS(get=lambda **kw: NS(name="wrong", display_name="other")))
    with pytest.raises(ValueError, match="does not match"):
        summary.resolve_submission(client, store, "submission", name="wrong")
    client.batches.get = lambda **kw: NS(name="right", display_name="submission")
    summary.resolve_submission(client, store, "submission", name="right")
    assert job["name"] == "right"
